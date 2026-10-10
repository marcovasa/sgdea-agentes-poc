#!/usr/bin/env python3
"""Convierte el Excel de incidencias de QA en issues de GitHub (idempotente).

Toma las filas cuyo Estado (columna H) sea distinto de "Aprobado" y crea un issue por
incidencia, con plantilla estándar y etiqueta `incidencia`. Si el issue de ese ID ya existe,
lo omite. Opcionalmente asigna de una vez un motor de agente.

Columnas esperadas (fila 1 = encabezados):
  A ID · B Título · C HU relacionada · D Severidad · E Pasos para reproducir ·
  F Resultado esperado · G Resultado obtenido · H Estado · I Criterios de aceptación · J Evidencias

Uso:
  pip install openpyxl
  python scripts/excel_a_issues.py incidencias/incidencias_qa_demo.xlsx --repo org/repo --dry-run
  python scripts/excel_a_issues.py incidencias/incidencias_qa_demo.xlsx --repo org/repo --agente ambos
"""

from __future__ import annotations

import argparse
import json
import subprocess  # nosec B404 - solo invoca gh con argumentos en lista, sin shell
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class Incidencia:
    id: str
    titulo: str
    hu: str
    severidad: str
    pasos: str
    esperado: str
    obtenido: str
    estado: str
    criterios: str
    evidencias: str

    @property
    def titulo_issue(self) -> str:
        return f"[{self.id}] {self.titulo}"

    def cuerpo(self) -> str:
        criterios = "\n".join(
            linea if linea.strip().startswith("- [") else f"- [ ] {linea.lstrip('- ').strip()}"
            for linea in self.criterios.splitlines()
            if linea.strip()
        )
        return (
            f"**ID QA:** {self.id}  \n**HU relacionada:** {self.hu}  \n"
            f"**Severidad:** {self.severidad}  \n**Estado en QA:** {self.estado}\n\n"
            f"### Pasos para reproducir\n{self.pasos}\n\n"
            f"### Resultado esperado\n{self.esperado}\n\n"
            f"### Resultado obtenido\n{self.obtenido}\n\n"
            f"### Criterios de aceptación\n{criterios}\n\n"
            f"### Evidencias\n{self.evidencias or '_Sin evidencias adjuntas._'}\n\n"
            "_Issue generado automáticamente desde el Excel de incidencias de QA._"
        )


def _texto(valor: object) -> str:
    return "" if valor is None else str(valor).strip()


def leer_excel(ruta: Path) -> list[Incidencia]:
    hoja = load_workbook(ruta, read_only=True, data_only=True).active
    incidencias = []
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        celdas = [_texto(c) for c in (list(fila) + [None] * 10)[:10]]
        if not celdas[0]:
            continue
        incidencias.append(Incidencia(*celdas))
    return incidencias


def gh(*args: str) -> str:
    resultado = subprocess.run(["gh", *args], check=True, capture_output=True, text=True)  # nosec B603 B607
    return resultado.stdout


def titulos_existentes(repo: str) -> set[str]:
    salida = gh("issue", "list", "--repo", repo, "--state", "all", "--limit", "1000",
                "--label", "incidencia", "--json", "title")
    return {i["title"] for i in json.loads(salida)}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("excel", type=Path)
    parser.add_argument("--repo", required=True, help="owner/repo")
    parser.add_argument("--agente", choices=["gemini", "claude", "ambos"],
                        help="Asigna un motor de agente a los issues creados")
    parser.add_argument("--dry-run", action="store_true", help="Solo muestra lo que haría")
    args = parser.parse_args()

    pendientes = [i for i in leer_excel(args.excel) if i.estado.lower() != "aprobado"]
    existentes = set() if args.dry_run else titulos_existentes(args.repo)
    print(f"{len(pendientes)} incidencias no aprobadas en {args.excel.name}")

    for inc in pendientes:
        if inc.titulo_issue in existentes:
            print(f"  = {inc.id} ya existe, se omite")
            continue
        etiquetas = ["incidencia"] + ([f"agente:{args.agente}"] if args.agente else [])
        if args.dry_run:
            print(f"  + {inc.titulo_issue}  [{', '.join(etiquetas)}]")
            continue
        url = gh("issue", "create", "--repo", args.repo, "--title", inc.titulo_issue,
                 "--body", inc.cuerpo(), "--label", "incidencia").strip()
        if args.agente:
            # Se etiqueta después de crear para que el evento "labeled" dispare el workflow del agente.
            gh("issue", "edit", url, "--add-label", f"agente:{args.agente}")
        print(f"  + {inc.id} → {url}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
